"""Bank statement parsers for CSV, Excel, and best-effort PDF imports."""
from __future__ import annotations

import csv
import hashlib
import io
import re
from datetime import date, datetime
from typing import Any, Dict, Iterable, List, Optional

from openpyxl import load_workbook

REQUIRED_COLUMNS = {"date", "amount", "direction", "merchant", "narration"}

DATE_KEYS = {"date", "txn date", "transaction date", "value date", "posted date"}
NARRATION_KEYS = {"narration", "description", "details", "particulars", "transaction details", "remarks"}
MERCHANT_KEYS = {"merchant", "payee", "beneficiary"}
AMOUNT_KEYS = {"amount", "transaction amount"}
DEBIT_KEYS = {"debit", "withdrawal", "withdrawals", "dr", "paid out"}
CREDIT_KEYS = {"credit", "deposit", "deposits", "cr", "paid in"}


def checksum(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def row_checksum(account_id: int, row: Dict[str, Any]) -> str:
    raw = "|".join(
        [
            str(account_id),
            row["txn_date"].isoformat(),
            str(round(float(row["amount"]), 2)),
            row["direction"],
            row.get("merchant") or "",
            row.get("narration") or "",
        ]
    )
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def parse_statement(filename: str, content: bytes) -> List[Dict[str, Any]]:
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "csv":
        return parse_csv(content)
    if suffix in {"xlsx", "xlsm"}:
        return parse_excel(content)
    if suffix == "pdf":
        return parse_pdf(content)
    raise ValueError("Supported statement formats: .csv, .xlsx, .xlsm, .pdf")


def parse_csv(content: bytes) -> List[Dict[str, Any]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise ValueError("CSV has no header row")

    normalized = {_clean_header(name): name for name in reader.fieldnames}
    if REQUIRED_COLUMNS.issubset(set(normalized)):
        return [_row_from_template(row) for row in reader]

    rows = []
    for raw in reader:
        row = _normalize_mapping({key: raw.get(original) for key, original in normalized.items()})
        if row:
            rows.append(row)
    if not rows:
        raise ValueError("No valid transactions found in CSV")
    return rows


def parse_excel(content: bytes) -> List[Dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        parsed = _parse_table_rows(rows)
        if parsed:
            return parsed
    raise ValueError("No valid transactions found in Excel workbook")


def parse_pdf(content: bytes) -> List[Dict[str, Any]]:
    try:
        import pdfplumber
    except ImportError as exc:
        raise ValueError("PDF support requires pdfplumber") from exc

    rows: List[Dict[str, Any]] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                parsed = _parse_table_rows(table)
                if parsed:
                    rows.extend(parsed)
            if rows:
                continue
            text = page.extract_text() or ""
            rows.extend(_parse_pdf_lines(text.splitlines()))
    if not rows:
        raise ValueError("No valid transactions found in PDF. Try Excel export if your bank provides it.")
    return rows


def _parse_table_rows(rows: Iterable[Iterable[Any]]) -> List[Dict[str, Any]]:
    row_list = [list(row) for row in rows if row and any(cell is not None and str(cell).strip() for cell in row)]
    for index, row in enumerate(row_list[:20]):
        headers = [_clean_header(cell) for cell in row]
        if _looks_like_header(headers):
            parsed = []
            for values in row_list[index + 1 :]:
                mapping = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
                normalized = _normalize_mapping(mapping)
                if normalized:
                    parsed.append(normalized)
            return parsed
    return []


def _looks_like_header(headers: List[str]) -> bool:
    header_set = set(headers)
    has_date = bool(header_set & DATE_KEYS)
    has_text = bool(header_set & NARRATION_KEYS or header_set & MERCHANT_KEYS)
    has_amount = bool(header_set & AMOUNT_KEYS or header_set & DEBIT_KEYS or header_set & CREDIT_KEYS)
    return has_date and has_text and has_amount


def _normalize_mapping(raw: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    txn_date = _first_date(raw, DATE_KEYS)
    narration = _first_text(raw, NARRATION_KEYS)
    merchant = _first_text(raw, MERCHANT_KEYS) or _merchant_from_narration(narration)

    amount = _first_number(raw, AMOUNT_KEYS)
    debit = _first_number(raw, DEBIT_KEYS)
    credit = _first_number(raw, CREDIT_KEYS)

    direction = None
    if debit and debit > 0:
        amount = debit
        direction = "debit"
    elif credit and credit > 0:
        amount = credit
        direction = "credit"
    elif amount is not None:
        direction = "debit" if amount < 0 else _direction_from_text(raw) or "debit"
        amount = abs(amount)

    if not txn_date or amount is None or not direction:
        return None

    return {
        "txn_date": txn_date,
        "amount": float(amount),
        "direction": direction,
        "merchant": merchant,
        "narration": narration,
    }


def _row_from_template(row: Dict[str, str]) -> Dict[str, Any]:
    txn_date = _parse_date(row["date"])
    direction = (row["direction"] or "").strip().lower()
    if direction not in {"debit", "credit"}:
        raise ValueError("direction must be debit or credit")
    return {
        "txn_date": txn_date,
        "amount": abs(float(_clean_number(row["amount"]))),
        "direction": direction,
        "merchant": row.get("merchant"),
        "narration": row.get("narration"),
    }


def _parse_pdf_lines(lines: Iterable[str]) -> List[Dict[str, Any]]:
    parsed = []
    pattern = re.compile(r"(?P<date>\d{1,2}[-/]\d{1,2}[-/]\d{2,4}).*?(?P<amount>-?\d[\d,]*\.?\d{0,2})\s*(?P<kind>dr|cr|debit|credit)?$", re.I)
    for line in lines:
        match = pattern.search(line.strip())
        if not match:
            continue
        txn_date = _parse_date(match.group("date"))
        amount = abs(float(_clean_number(match.group("amount"))))
        kind = (match.group("kind") or "debit").lower()
        direction = "credit" if kind in {"cr", "credit"} else "debit"
        narration = line[: match.start("amount")].strip()
        parsed.append({
            "txn_date": txn_date,
            "amount": amount,
            "direction": direction,
            "merchant": _merchant_from_narration(narration),
            "narration": narration,
        })
    return parsed


def _clean_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _first_text(raw: Dict[str, Any], keys: set[str]) -> Optional[str]:
    for key in keys:
        value = raw.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return None


def _first_date(raw: Dict[str, Any], keys: set[str]) -> Optional[date]:
    for key in keys:
        value = raw.get(key)
        if value:
            parsed = _parse_date(value)
            if parsed:
                return parsed
    return None


def _first_number(raw: Dict[str, Any], keys: set[str]) -> Optional[float]:
    for key in keys:
        value = raw.get(key)
        if value is None or str(value).strip() == "":
            continue
        try:
            return float(_clean_number(value))
        except ValueError:
            continue
    return None


def _parse_date(value: Any) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _clean_number(value: Any) -> str:
    text = str(value).strip().replace(",", "")
    text = re.sub(r"\b(INR|Rs\.?|CR|DR)\b", "", text, flags=re.I).strip()
    if text in {"-", "--"}:
        return ""
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    return text


def _direction_from_text(raw: Dict[str, Any]) -> Optional[str]:
    haystack = " ".join(str(value).lower() for value in raw.values() if value is not None)
    if re.search(r"\b(cr|credit)\b", haystack):
        return "credit"
    if re.search(r"\b(dr|debit)\b", haystack):
        return "debit"
    return None


def _merchant_from_narration(narration: Optional[str]) -> Optional[str]:
    if not narration:
        return None
    text = re.sub(r"\s+", " ", narration).strip()
    text = re.sub(r"^(upi|neft|imps|rtgs|pos|atm|card|transfer)[-/:\s]+", "", text, flags=re.I)
    return text[:80] or None
